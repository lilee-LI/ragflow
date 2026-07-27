#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
import json
import os
import re
from io import BytesIO
from types import SimpleNamespace

import numpy as np
import psycopg2
import pytest
from psycopg2 import sql


pytestmark = [pytest.mark.gaussdb_integration, pytest.mark.gaussdb_both]

KB_ID = "abcdefabcdefabcdefabcdefabcdefab"


class FakeSQLChatModel:
    """Deterministic LLM boundary; the adapter and retriever remain real."""

    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    async def async_chat(self, system_prompt, messages, params, **kwargs):
        self.calls.append((system_prompt, messages, params, kwargs))
        if not self.responses:
            raise AssertionError("FakeSQLChatModel was called more times than specified")
        return self.responses.pop(0)


class FakeEmbeddingModel:
    def encode_queries(self, _text):
        return [0.1, 0.2, 0.3, 0.4], 4

    def encode(self, texts):
        return np.asarray([[0.1, 0.2, 0.3, 0.4] for _ in texts]), len(texts)


class RecordingDealer:
    """A transparent recorder around a real Dealer instance."""

    def __init__(self, dealer):
        self._dealer = dealer
        self.retrieval_calls = []
        self.sqls = []

    def __getattr__(self, name):
        return getattr(self._dealer, name)

    async def retrieval(self, *args, **kwargs):
        self.retrieval_calls.append((args, kwargs))
        return await self._dealer.retrieval(*args, **kwargs)

    def sql_retrieval(self, sql, fetch_size=128, format="json"):
        self.sqls.append(sql)
        return self._dealer.sql_retrieval(sql, fetch_size=fetch_size, format=format)


def _use_sql_table(gaussdb_env):
    table = f"ragflow_{gaussdb_env['table_prefix']}"
    gaussdb_env["created_tables"].add(table)
    return table


def _row(row_id, kb_id, doc_id, doc_name, amount, status="active", vector=None, content=None):
    content_text = f"{status} amount {amount}" if content is None else content
    return {
        "id": row_id,
        "kb_id": kb_id,
        "doc_id": doc_id,
        "docnm_kwd": doc_name,
        "content_ltks": content_text,
        "content_with_weight": content_text,
        "chunk_data": {"amount": amount, "status": status},
        "q_4_vec": vector or [0.1, 0.2, 0.3, 0.4],
    }


def _independent_connection():
    return psycopg2.connect(
        host=os.environ["GAUSSDB_HOST"],
        port=int(os.environ.get("GAUSSDB_PORT", "19995")),
        dbname=os.environ["GAUSSDB_DATABASE"],
        user=os.environ["GAUSSDB_USER"],
        password=os.environ["GAUSSDB_PASSWORD"],
        options="-c default_transaction_read_only=off",
    )


def _seed_rows_independently(table, rows):
    expected_keys = []
    with _independent_connection() as independent:
        with independent.cursor() as cur:
            for source_row in rows:
                row = dict(source_row)
                vector_columns = [column for column in row if re.fullmatch(r"q_(\d+)_vec", column)]
                if not vector_columns:
                    row["q_4_vec"] = [0.0, 0.0, 0.0, 0.0]
                    row["q_4_vec_valid"] = False
                    vector_columns = ["q_4_vec"]
                else:
                    for vector_column in vector_columns:
                        dim = re.fullmatch(r"q_(\d+)_vec", vector_column).group(1)
                        row[f"q_{dim}_vec_valid"] = True

                columns = list(row)
                placeholders = []
                params = []
                for column in columns:
                    value = row[column]
                    vector_match = re.fullmatch(r"q_(\d+)_vec", column)
                    if vector_match:
                        placeholders.append(sql.SQL("%s::floatvector({})").format(sql.SQL(vector_match.group(1))))
                        params.append("[" + ",".join(str(item) for item in value) + "]")
                    elif isinstance(value, (dict, list)):
                        placeholders.append(sql.SQL("%s::jsonb"))
                        params.append(json.dumps(value, ensure_ascii=False))
                    else:
                        placeholders.append(sql.Placeholder())
                        params.append(value)

                cur.execute(
                    sql.SQL("INSERT INTO {}.{} ({}) VALUES ({})").format(
                        sql.Identifier(os.environ["GAUSSDB_SCHEMA"]),
                        sql.Identifier(table),
                        sql.SQL(", ").join(map(sql.Identifier, columns)),
                        sql.SQL(", ").join(placeholders),
                    ),
                    params,
                )
                expected_keys.append((row["id"], row["kb_id"]))
        independent.commit()

        with independent.cursor() as cur:
            cur.execute(
                sql.SQL("SELECT id, kb_id FROM {}.{} ORDER BY id, kb_id").format(
                    sql.Identifier(os.environ["GAUSSDB_SCHEMA"]),
                    sql.Identifier(table),
                )
            )
            assert cur.fetchall() == sorted(expected_keys)


def _create_and_insert(conn, table, kb_id, rows):
    assert conn.create_idx(table, kb_id, 4) is True
    _seed_rows_independently(table, rows)


def _xlsx_bytes(rows):
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    binary = BytesIO()
    workbook.save(binary)
    return binary.getvalue()


def _read_kb_parser_config(kb_id):
    from api.db.db_models import DB

    with DB.connection_context():
        cursor = DB.execute_sql("SELECT parser_id, parser_config FROM knowledgebase WHERE id = %s", (kb_id,))
        row = cursor.fetchone()
    assert row is not None
    parser_config = row[1]
    if isinstance(parser_config, bytes):
        parser_config = parser_config.decode("utf-8")
    if isinstance(parser_config, str):
        parser_config = json.loads(parser_config)
    return row[0], parser_config


def _configure_use_sql(monkeypatch, conn, dealer):
    from common import settings

    monkeypatch.setattr(settings, "docStoreConn", conn, raising=False)
    monkeypatch.setattr(settings, "retriever", dealer, raising=False)


def _dialog_service():
    from api.db.services import dialog_service

    return dialog_service


def test_tc_sql_006_table_chunk_persists_real_field_map(
    gaussdb_env,
    ragflow_kb_context,
    table_name,
    monkeypatch,
):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from common import settings
    from rag.app.table import chunk as table_chunk
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    monkeypatch.setattr(settings, "docStoreConn", conn, raising=False)
    kb_id = ragflow_kb_context["kb_id"]
    assert KnowledgebaseService.update_by_id(kb_id, {"parser_id": "table"}) == 1
    chunks = table_chunk(
        "sales.xlsx",
        _xlsx_bytes([["amount", "status"], [100, "A"], [200, "active"]]),
        0,
        10,
        "English",
        lambda *_args: None,
        kb_id=kb_id,
        parser_config={"table_column_mode": "manual", "table_column_roles": {"amount": "metadata", "status": "metadata"}},
    )

    assert [chunk["chunk_data"] for chunk in chunks] == [{"amount": 100, "status": "A"}, {"amount": 200, "status": "active"}]
    parser_id, parser_config = _read_kb_parser_config(kb_id)
    assert parser_id == "table"
    assert parser_config["field_map"] == {"amount": "amount", "status": "status"}
    assert parser_config["table_column_names"] == ["amount", "status"]
    table = table_name(gaussdb_env, "sql_006")
    _create_and_insert(
        conn,
        table,
        kb_id,
        [
            {
                "id": "chunk-sql-006",
                "kb_id": kb_id,
                "doc_id": "doc-sql-006",
                "docnm_kwd": "sales.xlsx",
                "chunk_data": chunks[0]["chunk_data"],
                "q_4_vec": [0.1, 0.2, 0.3, 0.4],
            }
        ],
    )
    amount_path = parser_config["field_map"]["amount"]
    assert conn.sql(f"SELECT chunk_data #>> '{{{amount_path}}}' AS amount FROM {table} WHERE kb_id = '{kb_id}'") == {"columns": [{"name": "amount", "type": "text"}], "rows": [["100"]]}


def test_tc_sql_007_table_chunk_data_reaches_real_jsonb_table(gaussdb_env, ragflow_kb_context, monkeypatch, table_name, gaussdb_admin_conn):
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from common import settings
    from rag.app.table import chunk as table_chunk
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    monkeypatch.setattr(settings, "docStoreConn", conn, raising=False)
    kb_id = ragflow_kb_context["kb_id"]
    assert KnowledgebaseService.update_by_id(kb_id, {"parser_id": "table"}) == 1
    chunks = table_chunk(
        "sales.xlsx",
        _xlsx_bytes([["amount", "status"], [100, "A"], [200, "active"]]),
        0,
        10,
        "English",
        lambda *_args: None,
        kb_id=kb_id,
        parser_config={"table_column_mode": "manual", "table_column_roles": {"amount": "metadata", "status": "metadata"}},
    )
    assert len(chunks) == 2
    parsed = next(chunk for chunk in chunks if chunk["chunk_data"]["amount"] == 200)
    assert parsed["chunk_data"] == {"amount": 200, "status": "active"}
    parser_id, parser_config = _read_kb_parser_config(kb_id)
    assert parser_id == "table"
    assert parser_config["field_map"] == {"amount": "amount", "status": "status"}

    table = table_name(gaussdb_env, "sql_007")
    row = dict(parsed)
    row.update({"id": "chunk-007", "kb_id": kb_id, "doc_id": "doc-007", "q_4_vec": [0.1, 0.2, 0.3, 0.4]})
    assert conn.create_idx(table, kb_id, 4) is True
    assert conn.insert([row], table, kb_id) == []
    with gaussdb_admin_conn.cursor() as cur:
        cur.execute(
            sql.SQL("SELECT chunk_data, q_4_vec_valid FROM {}.{} WHERE id = %s AND kb_id = %s").format(
                sql.Identifier(gaussdb_env["schema"]),
                sql.Identifier(table),
            ),
            ["chunk-007", kb_id],
        )
        stored_chunk_data, vector_valid = cur.fetchone()
        assert stored_chunk_data == {"amount": 200, "status": "active"}
        assert vector_valid is True
        cur.execute(
            """
            SELECT data_type
              FROM information_schema.columns
             WHERE table_schema = %s AND table_name = %s AND column_name = 'chunk_data'
            """,
            [gaussdb_env["schema"], table],
        )
        assert cur.fetchone() == ("jsonb",)


def test_tc_sql_411_jsonb_missing_null_and_empty_values(gaussdb_env, table_name):
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = table_name(gaussdb_env, "sql_411")
    _create_and_insert(
        conn,
        table,
        KB_ID,
        [
            {"id": "c1", "kb_id": KB_ID, "doc_id": "missing-key", "docnm_kwd": "doc.xlsx", "chunk_data": {"other": 1}},
            {"id": "c2", "kb_id": KB_ID, "doc_id": "json-null", "docnm_kwd": "doc.xlsx", "chunk_data": {"status": None}},
            {"id": "c3", "kb_id": KB_ID, "doc_id": "empty-string", "docnm_kwd": "doc.xlsx", "chunk_data": {"status": ""}},
        ],
    )

    null_rows = conn.sql(f"SELECT doc_id FROM {table} WHERE (chunk_data #>> '{{status}}') IS NULL AND kb_id = '{KB_ID}' ORDER BY doc_id LIMIT 128")
    not_null_rows = conn.sql(f"SELECT doc_id FROM {table} WHERE (chunk_data #>> '{{status}}') IS NOT NULL AND kb_id = '{KB_ID}' ORDER BY doc_id LIMIT 128")
    assert null_rows == {"columns": [{"name": "doc_id", "type": "text"}], "rows": [["json-null"], ["missing-key"]]}
    assert not_null_rows == {"columns": [{"name": "doc_id", "type": "text"}], "rows": [["empty-string"]]}


def test_tc_sql_501_sql_uses_real_validator_and_returns_rows(gaussdb_env, table_name):
    from common.doc_store.gaussdb_conn_base import UnsafeGaussDBSQL
    from rag.utils.gaussdb_conn import GaussDBConnection
    from rag.utils.gaussdb_text_to_sql import build_validator

    conn = GaussDBConnection()
    table = table_name(gaussdb_env, "sql_501")
    other_tenant_table = table_name(gaussdb_env, "sql501b")
    _create_and_insert(
        conn,
        table,
        KB_ID,
        [_row("c1", KB_ID, "d1", "doc.xlsx", 100), _row("c2", KB_ID, "d2", "doc.xlsx", 200)],
    )
    _create_and_insert(
        conn,
        other_tenant_table,
        KB_ID,
        [_row("sentinel", KB_ID, "other-tenant-secret", "other.xlsx", 999)],
    )
    result = conn.sql(f"SELECT doc_id FROM {table} WHERE kb_id = '{KB_ID}' ORDER BY doc_id", fetch_size=1)

    assert result == {"columns": [{"name": "doc_id", "type": "text"}], "rows": [["d1"]]}
    forbidden_table = f"forbidden_{gaussdb_env['table_prefix']}"
    with pytest.raises(UnsafeGaussDBSQL, match=rf"table {re.escape(forbidden_table)} is not allowed"):
        conn.sql(f"SELECT doc_id FROM {forbidden_table} WHERE kb_id = '{KB_ID}'")
    for qualified_table in (f"{gaussdb_env['schema']}.{table}", "pg_catalog.pg_class"):
        with pytest.raises(UnsafeGaussDBSQL, match="cross-schema SQL is not allowed"):
            conn.sql(f"SELECT doc_id FROM {qualified_table} WHERE kb_id = '{KB_ID}'")
    assert conn.sql(f"SELECT doc_id FROM {other_tenant_table} WHERE kb_id = '{KB_ID}'")["rows"] == [["other-tenant-secret"]]
    tenant_validator = build_validator(table, [KB_ID], {"amount": "number"})
    with pytest.raises(UnsafeGaussDBSQL, match=rf"table {re.escape(other_tenant_table)} is not allowed"):
        tenant_validator.validate_and_patch(f"SELECT doc_id FROM {other_tenant_table} WHERE kb_id = '{KB_ID}'")
    assert conn.sql(f"SELECT doc_id FROM {other_tenant_table} WHERE kb_id = '{KB_ID}'")["rows"] == [["other-tenant-secret"]]


def _show_statement_timeout(conn):
    raw_conn = conn.pool.get_conn()
    cursor = raw_conn.cursor()
    try:
        cursor.execute("SHOW statement_timeout")
        return cursor.fetchone()[0]
    finally:
        cursor.close()
        conn.pool.put_conn(raw_conn)


def test_tc_sql_504_sql_statement_timeout_is_reset_after_checkout(
    gaussdb_env,
    table_name,
    monkeypatch,
    gaussdb_admin_conn,
):
    import rag.utils.gaussdb_conn as gaussdb_conn_module

    monkeypatch.setattr(gaussdb_conn_module, "SQL_QUERY_TIMEOUT_MS", 100)
    conn = gaussdb_conn_module.GaussDBConnection()
    table = table_name(gaussdb_env, "sql_504")
    _create_and_insert(
        conn,
        table,
        KB_ID,
        [_row("c000", KB_ID, "d000", "doc.xlsx", 0)],
    )
    before = _show_statement_timeout(conn)
    with gaussdb_admin_conn.cursor() as cur:
        cur.execute(sql.SQL("LOCK TABLE {} IN ACCESS EXCLUSIVE MODE").format(sql.Identifier(gaussdb_env["schema"], table)))
    try:
        with pytest.raises(psycopg2.Error) as exc_info:
            conn.sql(f"SELECT COUNT(*) AS total FROM {table} WHERE kb_id = '{KB_ID}'")
    finally:
        gaussdb_admin_conn.rollback()
    assert exc_info.value.pgcode == "57014"
    assert "statement timeout" in str(exc_info.value).lower() or "canceling statement" in str(exc_info.value).lower()
    after = _show_statement_timeout(conn)

    assert after == before
    assert conn.sql(f"SELECT doc_id FROM {table} WHERE kb_id = '{KB_ID}' AND doc_id = 'd000'")["rows"] == [["d000"]]


def test_tc_sql_505_sql_returns_columns_and_rows(gaussdb_env, table_name):
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = table_name(gaussdb_env, "sql_505")
    _create_and_insert(
        conn,
        table,
        KB_ID,
        [_row("c1", KB_ID, "d1", "doc.xlsx", 100), _row("c2", KB_ID, "d2", "doc.xlsx", 200)],
    )
    result = conn.sql(f"SELECT doc_id, docnm_kwd FROM {table} WHERE kb_id = '{KB_ID}' ORDER BY doc_id LIMIT 128")

    assert result["columns"] == [{"name": "doc_id", "type": "text"}, {"name": "docnm_kwd", "type": "text"}]
    assert result["rows"] == [["d1", "doc.xlsx"], ["d2", "doc.xlsx"]]


def test_tc_sql_507_sql_preserves_columns_for_empty_rows(gaussdb_env, table_name):
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = table_name(gaussdb_env, "sql_507")
    _create_and_insert(conn, table, KB_ID, [_row("c1", KB_ID, "d1", "doc.xlsx", 100)])
    result = conn.sql(f"SELECT doc_id FROM {table} WHERE kb_id = '{KB_ID}' AND doc_id = 'missing-doc'")

    assert result == {"columns": [{"name": "doc_id", "type": "text"}], "rows": []}


def test_tc_sql_508_sql_returns_real_markdown(gaussdb_env, table_name):
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = table_name(gaussdb_env, "sql_508")
    _create_and_insert(
        conn,
        table,
        KB_ID,
        [_row("c1", KB_ID, "d1", "sales.xlsx", 100), _row("c2", KB_ID, "d2", "sales.xlsx", 200)],
    )
    result = conn.sql(
        f"SELECT doc_id, chunk_data #>> '{{amount}}' AS amount FROM {table} WHERE kb_id = '{KB_ID}' ORDER BY doc_id LIMIT 128",
        format="markdown",
    )

    assert result["markdown"] == "|doc_id|amount|\n|---|---|\n|d1|100|\n|d2|200|"


@pytest.mark.asyncio
async def test_tc_sql_509_adapter_database_error_and_use_sql_retry_are_real(gaussdb_env, monkeypatch):
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    dialog_service = _dialog_service()
    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    _create_and_insert(conn, table, KB_ID, [_row("c1", KB_ID, "d1", "doc.xlsx", 100)])
    invalid_sql = f"SELECT CAST(chunk_data #>> '{{status}}' AS INTEGER) AS amount, doc_id, docnm_kwd FROM {table} WHERE kb_id = '{KB_ID}'"
    with pytest.raises(psycopg2.Error) as exc_info:
        conn.sql(invalid_sql)
    assert exc_info.value.pgcode in {"22P02", "22018"}
    assert conn.sql(f"SELECT doc_id FROM {table} WHERE kb_id = '{KB_ID}'")["rows"] == [["d1"]]

    dealer = RecordingDealer(Dealer(conn))
    _configure_use_sql(monkeypatch, conn, dealer)
    valid_sql = f"SELECT doc_id, docnm_kwd, chunk_data #>> '{{amount}}' AS amount FROM {table} ORDER BY doc_id"
    chat = FakeSQLChatModel([invalid_sql, valid_sql])
    result = await dialog_service.use_sql(
        "show docs",
        {"amount": "amount", "status": "string"},
        gaussdb_env["table_prefix"],
        chat,
        quota=False,
        kb_ids=[KB_ID],
    )

    assert len(chat.calls) == 2
    assert "invalid" in str(chat.calls[1]).lower() or "integer" in str(chat.calls[1]).lower()
    assert len(dealer.sqls) == 2
    assert re.search(
        r"CAST\(chunk_data\s*#>>\s*'\{status\}'\s+AS\s+(?:INT|INTEGER)\)",
        dealer.sqls[0],
        flags=re.IGNORECASE,
    )
    executed_sql = dealer.sqls[-1]
    assert "doc_id" in executed_sql and "docnm_kwd" in executed_sql
    assert "chunk_data #>> '{amount}' AS amount" in executed_sql
    assert f"kb_id = '{KB_ID}'" in executed_sql
    assert "LIMIT" in executed_sql.upper()
    assert result["answer"] == "|amount|Source|\n|------|------|\n|100| ##0$$|"
    assert result["reference"]["chunks"] == [{"doc_id": "d1", "docnm_kwd": "doc.xlsx", "kb_id": KB_ID}]
    assert result["reference"]["doc_aggs"] == [{"doc_id": "d1", "doc_name": "doc.xlsx", "count": 1}]


async def _aggregate_result(gaussdb_env, monkeypatch, rows, sql, *, recording=False, field_map=None):
    from api.db.services import dialog_service
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    _create_and_insert(conn, table, KB_ID, rows)
    dealer = Dealer(conn)
    if recording:
        dealer = RecordingDealer(dealer)
    _configure_use_sql(monkeypatch, conn, dealer)
    result = await dialog_service.use_sql(
        "What is the total amount?",
        field_map or {"amount": "amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([sql]),
        quota=False,
        kb_ids=[KB_ID],
    )
    return result, dealer


@pytest.mark.asyncio
async def test_tc_sql_701_aggregate_sql_fetches_real_source_chunks(gaussdb_env, monkeypatch):
    from rag.utils.gaussdb_conn import GaussDBConnection
    from rag.utils.gaussdb_text_to_sql import build_aggregate_source_sql

    table = f"ragflow_{gaussdb_env['table_prefix']}"
    second_kb_id = "bbbbbbbbbbbbbbbbbbbbbbbbbbbbbbbb"
    sql = f"SELECT SUM(CAST(chunk_data #>> '{{amount}}' AS DOUBLE PRECISION)) AS total FROM {table} WHERE chunk_data #>> '{{status}}' = 'active' LIMIT 3"
    result, dealer = await _aggregate_result(
        gaussdb_env,
        monkeypatch,
        [
            _row("c1", KB_ID, "d1", "sales.xlsx", 200),
            _row("c2", KB_ID, "d1", "sales.xlsx", 150),
            _row("c3", KB_ID, "d2", "sales.xlsx", 100),
            _row("c4", KB_ID, "d3", "sales.xlsx", 50),
            _row("c5", second_kb_id, "d4", "other.xlsx", 999),
        ],
        sql,
        recording=True,
        field_map={"amount": "number", "status": "string"},
    )

    assert result["answer"] == "|total|\n|------\n|500.0|"
    assert {chunk["doc_id"] for chunk in result["reference"]["chunks"]} == {"d1", "d2", "d3"}
    assert {chunk["kb_id"] for chunk in result["reference"]["chunks"]} == {KB_ID}
    assert sorted(result["reference"]["doc_aggs"], key=lambda item: item["doc_id"]) == [
        {"doc_id": "d1", "doc_name": "sales.xlsx", "count": 2},
        {"doc_id": "d2", "doc_name": "sales.xlsx", "count": 1},
        {"doc_id": "d3", "doc_name": "sales.xlsx", "count": 1},
    ]
    assert len(dealer.sqls) == 2
    assert dealer.sqls[0].endswith("LIMIT 3")
    scoped_source = dealer.sqls[1]
    assert scoped_source.startswith("SELECT doc_id, docnm_kwd FROM")
    assert "chunk_data #>> '{status}' = 'active'" in scoped_source
    assert f"kb_id = '{KB_ID}'" in scoped_source
    assert scoped_source.endswith("LIMIT 128")

    multi_kb_aggregate = f"SELECT SUM(CAST(chunk_data #>> '{{amount}}' AS DOUBLE PRECISION)) AS total FROM {table} WHERE kb_id IN ('{KB_ID}', '{second_kb_id}')"
    multi_kb_source = build_aggregate_source_sql(multi_kb_aggregate, "docnm_kwd", include_kb_id=True)
    assert multi_kb_source.startswith("SELECT doc_id, docnm_kwd, kb_id FROM")
    multi_kb_rows = GaussDBConnection().sql(multi_kb_source)["rows"]
    assert sorted(multi_kb_rows) == [
        ["d1", "sales.xlsx", KB_ID],
        ["d1", "sales.xlsx", KB_ID],
        ["d2", "sales.xlsx", KB_ID],
        ["d3", "sales.xlsx", KB_ID],
        ["d4", "other.xlsx", second_kb_id],
    ]

    having_aggregate = f"SELECT SUM(CAST(chunk_data #>> '{{amount}}' AS DOUBLE PRECISION)) AS total FROM {table} HAVING SUM(CAST(chunk_data #>> '{{amount}}' AS DOUBLE PRECISION)) > 0"
    with pytest.raises(ValueError, match="cannot preserve HAVING safely"):
        build_aggregate_source_sql(having_aggregate, "docnm_kwd", include_kb_id=False)


@pytest.mark.asyncio
async def test_tc_sql_702_aggregate_source_lookup_is_real_validator_scoped(gaussdb_env, monkeypatch):
    table = f"ragflow_{gaussdb_env['table_prefix']}"
    sql = f"SELECT SUM(CAST(chunk_data #>> '{{amount}}' AS DOUBLE PRECISION)) AS total FROM {table}"
    result, dealer = await _aggregate_result(
        gaussdb_env,
        monkeypatch,
        [_row("c1", KB_ID, "d1", "sales.xlsx", 200)],
        sql,
        recording=True,
    )

    assert result["answer"] == "|total|\n|------\n|200.0|"
    assert result["reference"]["chunks"] == [{"doc_id": "d1", "docnm_kwd": "sales.xlsx", "kb_id": KB_ID}]
    assert result["reference"]["doc_aggs"] == [{"doc_id": "d1", "doc_name": "sales.xlsx", "count": 1}]
    assert len(dealer.sqls) == 2
    assert f"kb_id = '{KB_ID}'" in dealer.sqls[1]
    assert "SELECT doc_id, docnm_kwd" in dealer.sqls[1]
    assert "LIMIT 128" in dealer.sqls[1].upper()


@pytest.mark.asyncio
async def test_tc_sql_704_aggregate_without_where_gets_scoped_source_lookup(gaussdb_env, monkeypatch):
    table = f"ragflow_{gaussdb_env['table_prefix']}"
    sql = f"SELECT SUM(CAST(chunk_data #>> '{{amount}}' AS DOUBLE PRECISION)) AS total FROM {table}"
    result, dealer = await _aggregate_result(
        gaussdb_env,
        monkeypatch,
        [_row("c1", KB_ID, "d1", "sales.xlsx", 200), _row("c2", KB_ID, "d2", "sales.xlsx", 150)],
        sql,
        recording=True,
    )

    assert result["answer"] == "|total|\n|------\n|350.0|"
    assert result["reference"]["chunks"] == [
        {"doc_id": "d1", "docnm_kwd": "sales.xlsx", "kb_id": KB_ID},
        {"doc_id": "d2", "docnm_kwd": "sales.xlsx", "kb_id": KB_ID},
    ]
    assert result["reference"]["doc_aggs"] == [
        {"doc_id": "d1", "doc_name": "sales.xlsx", "count": 1},
        {"doc_id": "d2", "doc_name": "sales.xlsx", "count": 1},
    ]
    assert "WHERE kb_id = '" + KB_ID + "'" in dealer.sqls[1]
    assert "LIMIT 128" in dealer.sqls[1].upper()
    assert {chunk["kb_id"] for chunk in result["reference"]["chunks"]} == {KB_ID}


@pytest.mark.asyncio
async def test_tc_sql_802_single_kb_reference_gets_kb_id_without_lookup(gaussdb_env, monkeypatch):
    from api.db.services import dialog_service
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    _create_and_insert(conn, table, KB_ID, [_row("c1", KB_ID, "d1", "doc.xlsx", 100)])
    dealer = RecordingDealer(Dealer(conn))
    _configure_use_sql(monkeypatch, conn, dealer)
    sql = f"SELECT doc_id, docnm_kwd, chunk_data #>> '{{amount}}' AS amount FROM {table} ORDER BY doc_id"
    result = await dialog_service.use_sql(
        "show amount",
        {"amount": "amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([sql]),
        quota=False,
        kb_ids=[KB_ID],
    )

    assert result["reference"]["chunks"] == [{"doc_id": "d1", "docnm_kwd": "doc.xlsx", "kb_id": KB_ID}]
    assert len(dealer.sqls) == 1


@pytest.mark.asyncio
async def test_tc_sql_901_use_sql_answer_is_markdown(gaussdb_env, monkeypatch):
    from api.db.services import dialog_service
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    rows = [_row("c1", KB_ID, "d1", "doc.xlsx", 100), _row("c2", KB_ID, "d2", "doc.xlsx", 200)]
    for row in rows:
        row["chunk_data"] = {"Amount": row["chunk_data"]["amount"]}
    _create_and_insert(conn, table, KB_ID, rows)
    _configure_use_sql(monkeypatch, conn, Dealer(conn))
    sql = f"SELECT doc_id, docnm_kwd, chunk_data #>> '{{Amount}}' AS amount FROM {table} ORDER BY doc_id"
    result = await dialog_service.use_sql(
        "show amount",
        {"amount": "Amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([sql]),
        quota=False,
        kb_ids=[KB_ID],
    )

    assert result["answer"].splitlines()[0] == "|Amount|Source|"
    assert "|100| ##0$$|" in result["answer"]
    assert "|200| ##1$$|" in result["answer"]
    assert [chunk["doc_id"] for chunk in result["reference"]["chunks"]] == ["d1", "d2"]
    assert all(chunk["kb_id"] == KB_ID for chunk in result["reference"]["chunks"])


@pytest.mark.asyncio
async def test_tc_sql_902_unmapped_exposed_column_keeps_original_name(gaussdb_env, monkeypatch):
    from api.db.services import dialog_service
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    _create_and_insert(conn, table, KB_ID, [_row("c1", KB_ID, "d1", "doc.xlsx", 100)])
    _configure_use_sql(monkeypatch, conn, Dealer(conn))
    sql = f"SELECT doc_id, docnm_kwd, chunk_data #>> '{{amount}}' AS unlabelled_amount FROM {table} ORDER BY doc_id"
    result = await dialog_service.use_sql(
        "show amount and status",
        {"amount": "amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([sql]),
        quota=False,
        kb_ids=[KB_ID],
    )

    assert result["answer"].splitlines()[0] == "|unlabelled_amount|Source|"


@pytest.mark.asyncio
async def test_tc_sql_903_reference_chunks_have_source_shape(gaussdb_env, monkeypatch):
    from api.db.services import dialog_service
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    _create_and_insert(conn, table, KB_ID, [_row("c1", KB_ID, "d1", "doc.xlsx", 100), _row("c2", KB_ID, "d2", "doc.xlsx", 200)])
    _configure_use_sql(monkeypatch, conn, Dealer(conn))
    sql = f"SELECT doc_id, docnm_kwd, chunk_data #>> '{{amount}}' AS amount FROM {table} ORDER BY doc_id"
    result = await dialog_service.use_sql(
        "show amount",
        {"amount": "amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([sql]),
        quota=False,
        kb_ids=[KB_ID],
    )

    assert result["reference"]["chunks"] == [
        {"doc_id": "d1", "docnm_kwd": "doc.xlsx", "kb_id": KB_ID},
        {"doc_id": "d2", "docnm_kwd": "doc.xlsx", "kb_id": KB_ID},
    ]


@pytest.mark.asyncio
async def test_tc_sql_904_reference_doc_aggs_count_real_rows(gaussdb_env, monkeypatch):
    from api.db.services import dialog_service
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    _create_and_insert(
        conn,
        table,
        KB_ID,
        [_row("c1", KB_ID, "d1", "doc.xlsx", 100), _row("c2", KB_ID, "d1", "doc.xlsx", 150), _row("c3", KB_ID, "d2", "sales.xlsx", 200)],
    )
    _configure_use_sql(monkeypatch, conn, Dealer(conn))
    sql = f"SELECT doc_id, docnm_kwd FROM {table} ORDER BY doc_id"
    result = await dialog_service.use_sql(
        "show docs",
        {"amount": "amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([sql]),
        quota=False,
        kb_ids=[KB_ID],
    )

    assert sorted(result["reference"]["doc_aggs"], key=lambda item: item["doc_id"]) == [
        {"doc_id": "d1", "doc_name": "doc.xlsx", "count": 2},
        {"doc_id": "d2", "doc_name": "sales.xlsx", "count": 1},
    ]
    empty_result = await dialog_service.use_sql(
        "show missing docs",
        {"amount": "amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([f"SELECT COUNT(*) AS total FROM {table} WHERE chunk_data #>> '{{amount}}' = 'missing'"]),
        quota=False,
        kb_ids=[KB_ID],
    )
    assert empty_result["reference"]["chunks"] == []
    assert empty_result["reference"]["doc_aggs"] == []


@pytest.mark.asyncio
async def test_tc_sql_905_use_sql_returns_answer_reference_and_prompt(gaussdb_env, monkeypatch):
    from api.db.services import dialog_service
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    conn = GaussDBConnection()
    table = _use_sql_table(gaussdb_env)
    _create_and_insert(
        conn,
        table,
        KB_ID,
        [_row("c1", KB_ID, "d1", "sales.xlsx", 200), _row("c2", KB_ID, "d1", "sales.xlsx", 150)],
    )
    _configure_use_sql(monkeypatch, conn, Dealer(conn))
    sql = f"SELECT SUM(CAST(chunk_data #>> '{{amount}}' AS DOUBLE PRECISION)) AS total FROM {table}"
    result = await dialog_service.use_sql(
        "sum amount",
        {"amount": "amount"},
        gaussdb_env["table_prefix"],
        FakeSQLChatModel([sql]),
        quota=False,
        kb_ids=[KB_ID],
    )

    assert set(result) == {"answer", "reference", "prompt"}
    assert result["answer"] == "|total|\n|------\n|350.0|"
    for rule in (
        "GaussDB A/ORA",
        "Use only the static JSONB path literals",
        "Do not use json_extract",
        "Return exactly one read-only SELECT statement",
        "Do not use SELECT *, window functions, or system functions",
        "Do not use DML or DDL",
    ):
        assert rule in result["prompt"]
    assert result["reference"]["chunks"] == [
        {"doc_id": "d1", "docnm_kwd": "sales.xlsx", "kb_id": KB_ID},
        {"doc_id": "d1", "docnm_kwd": "sales.xlsx", "kb_id": KB_ID},
    ]
    assert result["reference"]["doc_aggs"] == [{"doc_id": "d1", "doc_name": "sales.xlsx", "count": 2}]


@pytest.mark.asyncio
async def test_tc_sql_1006_async_chat_falls_back_to_real_dealer_retrieval(
    gaussdb_env,
    ragflow_kb_context,
    monkeypatch,
    caplog,
    gaussdb_admin_conn,
):
    import logging

    from api.db.db_models import DB
    from api.db.services.document_service import DocumentService
    from api.db.services.knowledgebase_service import KnowledgebaseService
    from common import settings
    from psycopg2 import sql
    from rag.prompts.generator import kb_prompt, message_fit_in
    from rag.nlp.search import Dealer
    from rag.utils.gaussdb_conn import GaussDBConnection

    dialog_service = _dialog_service()
    monkeypatch.setattr(dialog_service, "KnowledgebaseService", KnowledgebaseService, raising=False)
    monkeypatch.setattr(dialog_service, "DocumentService", DocumentService, raising=False)
    monkeypatch.setattr(dialog_service, "TenantLangfuseService", SimpleNamespace(filter_by_tenant=lambda **_kwargs: None), raising=False)
    monkeypatch.setattr(dialog_service, "kb_prompt", kb_prompt, raising=False)
    monkeypatch.setattr(dialog_service, "message_fit_in", message_fit_in, raising=False)
    tenant_id = gaussdb_env["table_prefix"]
    kb_id = ragflow_kb_context["kb_id"]
    doc_id = "d-sql-1006"
    table = _use_sql_table(gaussdb_env)
    conn = GaussDBConnection()
    _create_and_insert(
        conn,
        table,
        kb_id,
        [
            _row(
                "c1",
                kb_id,
                doc_id,
                "fallback.csv",
                100,
                vector=[0.1, 0.2, 0.3, 0.4],
                content="fallback-gaussdb-query",
            )
        ],
    )
    events = None
    try:
        KnowledgebaseService.update_parser_config(kb_id, {"field_map": {"amount": "amount"}})
        DocumentService.save(
            id=doc_id,
            kb_id=kb_id,
            parser_id="naive",
            type="txt",
            created_by=tenant_id,
            name="fallback.csv",
            suffix="txt",
        )

        chat = FakeSQLChatModel([f"SELECT * FROM {table}", f"SELECT * FROM {table}", "fallback answer"])
        dealer = RecordingDealer(Dealer(conn))
        real_kbs = KnowledgebaseService.get_by_ids([kb_id])
        assert [kb.id for kb in real_kbs] == [kb_id]
        monkeypatch.setattr(settings, "docStoreConn", conn, raising=False)
        monkeypatch.setattr(settings, "retriever", dealer, raising=False)
        monkeypatch.setattr(
            dialog_service,
            "get_tenant_default_model_by_type",
            lambda *_args, **_kwargs: {"model_type": "chat", "max_tokens": 8192},
        )
        monkeypatch.setattr(
            dialog_service,
            "get_models",
            lambda *_args, **_kwargs: (real_kbs, FakeEmbeddingModel(), None, chat, None),
        )
        monkeypatch.setattr(dialog_service, "tts", lambda *_args, **_kwargs: None, raising=False)
        dialog = SimpleNamespace(
            tenant_id=tenant_id,
            kb_ids=[kb_id],
            prompt_config={"quote": True, "parameters": [{"key": "knowledge", "optional": False}], "system": "Use {knowledge}"},
            llm_id=None,
            tenant_llm_id=None,
            llm_setting={},
            top_n=6,
            top_k=10,
            similarity_threshold=0.0,
            vector_similarity_weight=1.0,
            meta_data_filter=None,
        )

        with caplog.at_level(logging.DEBUG):
            events = [
                item
                async for item in dialog_service.async_chat(
                    dialog,
                    [{"role": "user", "content": "fallback-gaussdb-query"}],
                    stream=False,
                )
            ]
            assert len(dealer.retrieval_calls) == 1
            retrieval_args, retrieval_kwargs = dealer.retrieval_calls[0]
            assert (retrieval_args[0] if retrieval_args else retrieval_kwargs["question"]) == "fallback-gaussdb-query"
            assert (retrieval_args[2] if retrieval_args else retrieval_kwargs["tenant_ids"]) == [tenant_id]
            assert (retrieval_args[3] if retrieval_args else retrieval_kwargs["kb_ids"]) == [kb_id]
            with gaussdb_admin_conn.cursor() as cur:
                cur.execute(
                    sql.SQL("SELECT COUNT(*) FROM {} WHERE kb_id = %s AND doc_id = %s").format(sql.Identifier(gaussdb_env["schema"], table)),
                    [kb_id, doc_id],
                )
                assert cur.fetchone() == (1,)
    finally:
        DocumentService.delete_by_id(doc_id)
        with DB.connection_context():
            cursor = DB.execute_sql("SELECT COUNT(*) FROM document WHERE id = %s", (doc_id,))
            assert cursor.fetchone() == (0,)

    assert events[-1]["answer"] == "fallback answer [ID:0]"
    assert events[-1]["reference"]["chunks"]
    assert events[-1]["reference"]["chunks"][0]["doc_id"] == doc_id
    assert events[-1]["reference"]["chunks"][0]["kb_id"] == kb_id
    assert events[-1]["reference"]["doc_aggs"] == [{"doc_id": doc_id, "doc_name": "fallback.csv", "count": 1}]
    assert "SQL failed or returned no results, falling back to vector search" in caplog.text
